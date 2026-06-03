"""
Generador de reportes Excel y PDF.
"""
import io
from datetime import datetime
from typing import Optional

from sqlalchemy.orm import Session

from backend.models import AttendanceRecord, Employee

from reportlab.platypus import Flowable
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
import os

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Draw a thin footer separator line
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        # Margin is 0.8*cm on sides (A4 width is 21.0cm)
        self.line(0.8*cm, 1.2*cm, 20.2*cm, 1.2*cm)
        
        # Left side text: system name / company
        cfg = _get_system_config()
        system_text = f"{cfg.get('system_name')} — {cfg.get('company_name')}"
        self.drawString(0.8*cm, 0.8*cm, system_text[:80])
        
        # Right side text: Page x of y
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(20.2*cm, 0.8*cm, page_text)
        self.restoreState()

class CircularAvatar(Flowable):
    def __init__(self, path_or_initials, size):
        Flowable.__init__(self)
        self.path = path_or_initials if (path_or_initials and os.path.exists(path_or_initials)) else None
        self.size = size
        self.width = size
        self.height = size

    def draw(self):
        self.canv.saveState()
        # Create circular clipping path
        path = self.canv.beginPath()
        path.circle(self.size / 2, self.size / 2, self.size / 2)
        self.canv.clipPath(path, stroke=0, fill=0)
        
        if self.path:
            self.canv.drawImage(self.path, 0, 0, width=self.size, height=self.size)
        else:
            # Draw placeholder avatar
            self.canv.setFillColor(colors.HexColor("#cbd5e1"))
            self.canv.circle(self.size / 2, self.size / 2, self.size / 2, fill=1, stroke=0)
            # Head
            self.canv.setFillColor(colors.white)
            self.canv.circle(self.size / 2, self.size * 0.62, self.size * 0.18, fill=1, stroke=0)
            # Body
            self.canv.circle(self.size / 2, self.size * 0.15, self.size * 0.32, fill=1, stroke=0)
            
        self.canv.restoreState()

def generate_excel_report(
    db: Session,
    employee_id: Optional[int] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    department_id: Optional[int] = None,
    position_id: Optional[int] = None,
    schedule_id: Optional[int] = None,
) -> bytes:
    """Genera un reporte Excel con los registros de asistencia filtrados."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    query = db.query(AttendanceRecord).join(
        Employee, AttendanceRecord.employee_id == Employee.id, isouter=True
    )
    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if position_id:
        query = query.filter(Employee.position_id == position_id)
    if schedule_id:
        query = query.filter(Employee.schedule_id == schedule_id)
    if date_from:
        query = query.filter(AttendanceRecord.event_time >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.event_time <= date_to)

    records = query.order_by(AttendanceRecord.event_time.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "Empleado", "Código", "Departamento", "Fecha", "Hora", "Tipo", "Método Auth", "Tardanza"]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Datos
    for i, r in enumerate(records, 1):
        emp = r.employee
        row = [
            i,
            emp.full_name if emp else "Desconocido",
            emp.employee_code if emp else "-",
            emp.department.name if emp and emp.department else "-",
            r.event_time.strftime("%d/%m/%Y"),
            r.event_time.strftime("%H:%M:%S"),
            "Entrada" if r.event_type == "entry" else "Salida",
            r.auth_method or "-",
            "Sí" if r.is_late else "No",
        ]
        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.alignment = center
            cell.border = border
            if r.is_late and col_idx == 9:
                cell.font = Font(color="CC0000", bold=True)

    # Anchos de columna
    col_widths = [5, 28, 12, 20, 14, 12, 10, 16, 10]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
def generate_pdf_report(
    db: Session,
    employee_id: Optional[int] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    department_id: Optional[int] = None,
    position_id: Optional[int] = None,
    schedule_id: Optional[int] = None,
) -> bytes:
    """Genera un reporte PDF con los registros de asistencia."""
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import cm

    query = db.query(AttendanceRecord).join(
        Employee, AttendanceRecord.employee_id == Employee.id, isouter=True
    )
    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if position_id:
        query = query.filter(Employee.position_id == position_id)
    if schedule_id:
        query = query.filter(Employee.schedule_id == schedule_id)
    if date_from:
        query = query.filter(AttendanceRecord.event_time >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.event_time <= date_to)

    records = query.order_by(AttendanceRecord.event_time.desc()).all()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#1e3a5f"))
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.grey)

    from backend.utils import get_local_now
    elements = [
        Paragraph("Reporte de Asistencia", title_style),
        Paragraph(f"Generado: {get_local_now().strftime('%d/%m/%Y %H:%M')} — Total registros: {len(records)}", sub_style),
        Spacer(1, 0.5*cm),
    ]

    data = [["#", "Empleado", "Código", "Departamento", "Fecha", "Hora", "Tipo", "Tardanza"]]
    for i, r in enumerate(records, 1):
        emp = r.employee
        data.append([
            str(i),
            emp.full_name if emp else "Desconocido",
            emp.employee_code if emp else "-",
            emp.department.name if emp and emp.department else "-",
            r.event_time.strftime("%d/%m/%Y"),
            r.event_time.strftime("%H:%M"),
            "Entrada" if r.event_type == "entry" else "Salida",
            "Sí" if r.is_late else "No",
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4f8")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()

def generate_consolidated_excel(
    db: Session,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    department_id: Optional[int] = None,
    position_id: Optional[int] = None,
    schedule_id: Optional[int] = None,
    columns: Optional[list] = None,
) -> bytes:
    """Genera un reporte consolidado en Excel con entrada y salida en la misma fila por día."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    query = db.query(AttendanceRecord).join(
        Employee, AttendanceRecord.employee_id == Employee.id, isouter=True
    )
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if position_id:
        query = query.filter(Employee.position_id == position_id)
    if schedule_id:
        query = query.filter(Employee.schedule_id == schedule_id)
    if date_from:
        query = query.filter(AttendanceRecord.event_time >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.event_time <= date_to)

    records = query.order_by(AttendanceRecord.event_time.asc()).all()

    # Agrupar por (empleado, fecha)
    data_map = defaultdict(lambda: {"entries": [], "exits": []})
    employee_cache = {}

    for r in records:
        if not r.employee_id:
            continue
        emp_id = r.employee_id
        if emp_id not in employee_cache:
            employee_cache[emp_id] = r.employee
            
        day = r.event_time.date()
        if r.event_type == "entry":
            data_map[(emp_id, day)]["entries"].append(r)
        else:
            data_map[(emp_id, day)]["exits"].append(r)

    # Ordenar por fecha desc, luego nombre empleado
    sorted_keys = sorted(
        data_map.keys(),
        key=lambda k: (k[1], employee_cache.get(k[0]).last_name if employee_cache.get(k[0]) else ""),
        reverse=True
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado Diario"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    all_cols = [
        ("period", "Fecha"),
        ("employee_code", "Código"),
        ("employee_name", "Empleado"),
        ("department", "Departamento"),
        ("punches", "Hora Entrada"),
        ("punches", "Hora Salida"),
        ("schedule", "Horario"),
        ("status", "Tardanza")
    ]

    active_cols = []
    for col_key, col_header in all_cols:
        if columns is None or col_key in columns:
            active_cols.append((col_key, col_header))

    headers = [h for _, h in active_cols]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    row_num = 2
    for key in sorted_keys:
        emp_id, day = key
        emp = employee_cache.get(emp_id)
        if not emp:
            continue
            
        group = data_map[key]
        
        # Entrada: primer registro de entrada del día
        entry_time_str = "-"
        is_late_str = "No"
        if group["entries"]:
            first_entry = min(group["entries"], key=lambda r: r.event_time)
            entry_time_str = first_entry.event_time.strftime("%H:%M:%S")
            is_late_str = "Sí" if first_entry.is_late else "No"
            
        # Salida: último registro de salida del día
        exit_time_str = "-"
        if group["exits"]:
            last_exit = max(group["exits"], key=lambda r: r.event_time)
            exit_time_str = last_exit.event_time.strftime("%H:%M:%S")

        full_row_data = {
            "period": day.strftime("%d/%m/%Y"),
            "employee_code": emp.employee_code,
            "employee_name": emp.full_name,
            "department": emp.department.name if emp.department else "-",
            "punches_in": entry_time_str,
            "punches_out": exit_time_str,
            "schedule": f"{emp.schedule.name} ({emp.schedule.work_start_time}-{emp.schedule.work_end_time})" if emp.schedule_id and emp.schedule else f"{emp.work_start_time} - {emp.work_end_time}",
            "status": is_late_str
        }

        row = []
        for col_key, col_header in active_cols:
            if col_header == "Hora Entrada":
                row.append(full_row_data["punches_in"])
            elif col_header == "Hora Salida":
                row.append(full_row_data["punches_out"])
            else:
                row.append(full_row_data[col_key])

        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = center
            cell.border = border
            header_name = headers[col_idx - 1]
            if header_name == "Tardanza" and is_late_str == "Sí":
                cell.font = Font(color="CC0000", bold=True)
        
        row_num += 1

    # Anchos de columna
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_attendance_excel(summaries: list, granularity: str, columns: Optional[list] = None) -> bytes:
    """Genera un reporte consolidado con horarios y granularidad en Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {granularity.capitalize()}"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if granularity == "daily":
        all_cols = [
            ("employee_name", "Empleado"),
            ("employee_code", "Código"),
            ("department", "Departamento"),
            ("period", "Fecha"),
            ("schedule", "Horario"),
            ("punches", "Entrada"),
            ("punches", "Salida Almuerzo"),
            ("punches", "Retorno Almuerzo"),
            ("punches", "Salida"),
            ("status", "Estado")
        ]
    else:
        all_cols = [
            ("employee_name", "Empleado"),
            ("employee_code", "Código"),
            ("department", "Departamento"),
            ("period", "Período"),
            ("schedule", "Horario"),
            ("status", "Asistente"),
            ("status", "Incidencias Tardanza"),
            ("status", "Punches Incompletos"),
            ("punches", "Eventos Totales")
        ]

    active_cols = []
    for col_key, col_header in all_cols:
        if columns is None or col_key in columns:
            active_cols.append((col_key, col_header))

    headers = [h for _, h in active_cols]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    row_num = 2
    for s in summaries:
        if granularity == "daily":
            def format_time(iso_str):
                if not iso_str:
                    return "-"
                try:
                    return datetime.fromisoformat(iso_str).strftime("%H:%M")
                except:
                    return "-"

            is_split = s["schedule_type"] == "split"
            entry_1 = format_time(s["punches"]["entry_1"])
            exit_1 = format_time(s["punches"]["exit_1"]) if is_split else "-"
            entry_2 = format_time(s["punches"]["entry_2"]) if is_split else "-"
            exit_2 = format_time(s["punches"]["exit_2"]) if is_split else format_time(s["punches"]["exit_1"])

            leave_map = {
                "vacation": "Vacaciones",
                "medical": "Incapacidad",
                "paid_leave": "Licencia Remunerada",
                "unpaid_leave": "Licencia No Remunerada",
                "suspension": "Suspensión"
            }
            status = "OK"
            if s.get("justification"):
                if s["justification"]["override_status"] == "present":
                    status = "Justificado"
                else:
                    status = "A Tiempo (Just.)"
            elif s.get("leave_type") and s["leave_type"] in leave_map:
                status = leave_map[s["leave_type"]]
            elif not s["is_present"]:
                status = "Ausente"
            elif s["missing_punches"]:
                status = "Incompleto"
            elif s["is_late"]:
                status = "Tardanza"

            try:
                formatted_date = datetime.strptime(s["date"], "%Y-%m-%d").strftime("%d/%m/%Y")
            except:
                formatted_date = s["date"]

            full_row_data = {
                "employee_name": s["employee_name"],
                "employee_code": s["employee_code"],
                "department": s["department"],
                "period": formatted_date,
                "schedule": "Partido" if is_split else ("Continuo" if s["schedule_type"] == "continuous" else "Sin Horario"),
                "punches_e1": entry_1,
                "punches_x1": exit_1,
                "punches_e2": entry_2,
                "punches_x2": exit_2,
                "status": status
            }

            row = []
            for col_key, col_header in active_cols:
                if col_header == "Entrada":
                    row.append(full_row_data["punches_e1"])
                elif col_header == "Salida Almuerzo":
                    row.append(full_row_data["punches_x1"])
                elif col_header == "Retorno Almuerzo":
                    row.append(full_row_data["punches_e2"])
                elif col_header == "Salida":
                    row.append(full_row_data["punches_x2"])
                else:
                    row.append(full_row_data[col_key])
        else:
            try:
                p_start = datetime.strptime(s["period_start"], "%Y-%m-%d").strftime("%d/%m/%Y")
                p_end = datetime.strptime(s["period_end"], "%Y-%m-%d").strftime("%d/%m/%Y")
                period_str = f"{p_start} - {p_end}"
            except:
                period_str = f"{s['period_start']} - {s['period_end']}"

            full_row_data = {
                "employee_name": s["employee_name"],
                "employee_code": s["employee_code"],
                "department": s["department"],
                "period": period_str,
                "schedule": "Partido" if s["schedule_type"] == "split" else ("Continuo" if s["schedule_type"] == "continuous" else "Sin Horario"),
                "status_p": "Sí" if s["is_present"] else "No",
                "status_l": "Sí" if s["is_late"] else "No",
                "status_m": "Sí" if s["missing_punches"] else "No",
                "punches": s["total_raw_events"]
            }

            row = []
            for col_key, col_header in active_cols:
                if col_header == "Asistente":
                    row.append(full_row_data["status_p"])
                elif col_header == "Incidencias Tardanza":
                    row.append(full_row_data["status_l"])
                elif col_header == "Punches Incompletos":
                    row.append(full_row_data["status_m"])
                else:
                    row.append(full_row_data[col_key])

        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = center
            cell.border = border
            header_name = headers[col_idx - 1]
            if granularity == "daily" and header_name == "Estado":
                if status == "Ausente":
                    cell.font = Font(color="FF3D00", bold=True)
                elif status == "Tardanza":
                    cell.font = Font(color="FFB300", bold=True)
                elif status == "Incompleto":
                    cell.font = Font(color="FFA000", bold=True)
                else:
                    cell.font = Font(color="00E676", bold=True)
            elif granularity != "daily":
                if header_name == "Incidencias Tardanza" and s["is_late"]:
                    cell.font = Font(color="FFB300", bold=True)
                elif header_name == "Punches Incompletos" and s["missing_punches"]:
                    cell.font = Font(color="FFA000", bold=True)

        row_num += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _get_system_config() -> dict:
    """Obtiene la configuración de marca del sistema desde la base de datos."""
    from backend.database import SessionLocal
    from backend.models import SystemConfig
    db = SessionLocal()
    try:
        cfg = db.query(SystemConfig).first()
        if cfg:
            return {
                "system_name": cfg.system_name,
                "company_name": cfg.company_name,
                "logo_path": cfg.logo_path,
                "primary_color": cfg.primary_color,
                "accent_color": cfg.accent_color,
                "work_days": cfg.work_days,
                "time_format": cfg.time_format,
                "entry_tolerance_minutes": cfg.entry_tolerance_minutes
            }
    except Exception:
        pass
    finally:
        db.close()
    
    return {
        "system_name": "Control de Asistencia",
        "company_name": "Hikvision DS-K1T323MBWX",
        "logo_path": None,
        "primary_color": "#1e3a5f",
        "accent_color": "#00e676",
        "work_days": "1,2,3,4,5",
        "time_format": "24h",
        "entry_tolerance_minutes": 10
    }


def _generate_matplotlib_charts(summaries: list, granularity: str) -> tuple:
    """Genera gráficos de distribución y tendencia utilizando matplotlib y los retorna como BytesIO."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import io
    from datetime import datetime

    cfg = _get_system_config()
    primary_color = cfg.get("primary_color") or "#1e3a5f"

    # Executive pastel colors
    color_ontime = "#4ade80"      # soft green
    color_late = "#fbbf24"        # soft orange/yellow
    color_incomplete = "#f97316"  # soft orange
    color_absent = "#f87171"      # soft red

    # 1. Gráfico Doughnut (Distribución de asistencia)
    ontime = 0
    late = 0
    absent = 0
    incomplete = 0
    for s in summaries:
        if not s.get("is_present", False):
            absent += 1
        elif s.get("missing_punches", False):
            incomplete += 1
        elif s.get("is_late", False):
            late += 1
        else:
            ontime += 1

    labels = []
    sizes = []
    colors_list = []
    if ontime > 0:
        labels.append(f"A tiempo ({ontime})")
        sizes.append(ontime)
        colors_list.append(color_ontime)
    if late > 0:
        labels.append(f"Tardanza ({late})")
        sizes.append(late)
        colors_list.append(color_late)
    if incomplete > 0:
        labels.append(f"Incompleto ({incomplete})")
        sizes.append(incomplete)
        colors_list.append(color_incomplete)
    if absent > 0:
        labels.append(f"Ausente ({absent})")
        sizes.append(absent)
        colors_list.append(color_absent)

    if not sizes:
        labels.append("Sin registros")
        sizes.append(1)
        colors_list.append("#e2e8f0")

    fig, ax = plt.subplots(figsize=(1.8, 1.8), dpi=150)
    wedges, texts, autotexts = ax.pie(
        sizes, labels=labels, autopct='%1.0f%%', startangle=90,
        colors=colors_list, textprops=dict(color="#1e293b", size=5.0, weight="bold"),
        wedgeprops=dict(width=0.35, edgecolor='white', linewidth=0.7)
    )
    plt.setp(autotexts, size=4.5, weight="bold", color="white")
    ax.set_title("Distribución de Estados", fontsize=6.5, weight="bold", color=primary_color, pad=8)
    fig.tight_layout()
    
    img_buf_1 = io.BytesIO()
    fig.savefig(img_buf_1, format='png', bbox_inches='tight', transparent=True)
    plt.close(fig)
    img_buf_1.seek(0)

    # 2. Gráfico de Tendencia (Horas de entrada)
    dates = []
    entry_minutes = []
    
    sorted_summaries = sorted(summaries, key=lambda x: x.get("date", ""))
    
    for s in sorted_summaries:
        if s.get("is_present", False) and s.get("punches") and s["punches"].get("entry_1"):
            try:
                dt_obj = datetime.fromisoformat(s["punches"]["entry_1"])
                min_past_midnight = dt_obj.hour * 60 + dt_obj.minute
                d_str = dt_obj.strftime("%d/%m")
                dates.append(d_str)
                entry_minutes.append(min_past_midnight)
            except Exception:
                pass

    dates = dates[-12:]
    entry_minutes = entry_minutes[-12:]

    fig2, ax2 = plt.subplots(figsize=(3.2, 1.8), dpi=150)
    if entry_minutes and granularity == "daily":
        ax2.plot(dates, entry_minutes, marker='o', color=primary_color, linewidth=1.2, markersize=2.5, label='Entrada')
        ax2.fill_between(dates, entry_minutes, color=primary_color, alpha=0.06)
        
        def format_min_to_hm(x, pos):
            h = int(x // 60)
            m = int(x % 60)
            return f"{h:02d}:{m:02d}"
        
        from matplotlib.ticker import FuncFormatter
        ax2.yaxis.set_major_formatter(FuncFormatter(format_min_to_hm))
        ax2.axhline(y=480, color=color_absent, linestyle='--', linewidth=0.75, alpha=0.8, label='Límite (08:00)')
        ax2.legend(fontsize=5.0, loc='upper right', framealpha=0.9, edgecolor='#cbd5e1')
    else:
        msg = 'Tendencia disponible en vista Diaria' if granularity != "daily" else 'Sin entradas registradas'
        ax2.text(0.5, 0.5, msg, horizontalalignment='center', verticalalignment='center', transform=ax2.transAxes, color="#64748b", fontsize=6.5)

    ax2.set_title("Puntualidad en Entradas (Últimos 12 Días)", fontsize=6.5, weight="bold", color=primary_color, pad=8)
    ax2.tick_params(axis='both', which='major', labelsize=5.5, labelcolor="#475569")
    for label in ax2.get_xticklabels():
        label.set_rotation(30)
    ax2.grid(True, linestyle='--', alpha=0.4, color="#e2e8f0")
    for spine in ['top', 'right']:
        ax2.spines[spine].set_visible(False)
    ax2.spines['left'].set_color('#e2e8f0')
    ax2.spines['bottom'].set_color('#e2e8f0')
    
    fig2.tight_layout()

    img_buf_2 = io.BytesIO()
    fig2.savefig(img_buf_2, format='png', bbox_inches='tight', transparent=True)
    plt.close(fig2)
    return img_buf_1, img_buf_2


def _get_employee_flowables(summaries: list, granularity: str, schedules_map: dict = None, columns: Optional[list] = None) -> list:
    """Genera la lista de flowables (elementos de ReportLab) para un empleado individual."""
    import os
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.units import cm
    from backend.utils import get_local_now

    cfg = _get_system_config()
    primary_color_hex = cfg.get("primary_color") or "#1e3a5f"
    accent_color_hex = cfg.get("accent_color") or "#00e676"
    
    first_record = summaries[0] if summaries else {}
    emp_id = first_record.get("employee_id")
    
    work_days = None
    if schedules_map is not None and emp_id in schedules_map:
        work_days = schedules_map[emp_id]
        
    photo_path = None
    if emp_id:
        from backend.database import SessionLocal
        from backend.models import Employee
        db = SessionLocal()
        try:
            emp = db.query(Employee).filter(Employee.id == emp_id).first()
            if emp:
                photo_path = emp.photo_path
                if work_days is None and emp.schedule and emp.schedule.work_days:
                    work_days = [int(x) for x in emp.schedule.work_days.split(",")]
        except Exception as e:
            print(f"Error cargando datos de empleado en reporte PDF: {e}")
        finally:
            db.close()
            
    if work_days is None:
        work_days = [int(x) for x in cfg.get("work_days", "1,2,3,4,5").split(",")]

    # Precargar horarios diarios para el empleado en el rango analizado
    daily_schedules_map = {}
    if emp_id and summaries:
        from backend.database import SessionLocal
        from backend.models import EmployeeDailySchedule
        db = SessionLocal()
        try:
            dates = []
            from datetime import datetime as dt_parser
            for s in summaries:
                try:
                    dates.append(dt_parser.strptime(s["date"], "%Y-%m-%d").date())
                except:
                    pass
            if dates:
                min_date = min(dates)
                max_date = max(dates)
                dailies = db.query(EmployeeDailySchedule).filter(
                    EmployeeDailySchedule.employee_id == emp_id,
                    EmployeeDailySchedule.date >= min_date,
                    EmployeeDailySchedule.date <= max_date
                ).all()
                for d in dailies:
                    daily_schedules_map[d.date] = d
        except Exception as e:
            print(f"Error cargando horarios diarios en reporte: {e}")
        finally:
            db.close()

    primary_color = colors.HexColor(primary_color_hex)
    accent_color = colors.HexColor(accent_color_hex)

    employee_name = first_record.get("employee_name", "Empleado")
    employee_code = first_record.get("employee_code", "-")
    department = first_record.get("department", "-")

    # Calcular KPIs
    total_days = len(summaries)
    present_days = sum(1 for s in summaries if s.get("is_present", False))
    
    # Calcular ausencias únicamente en días laborales configurados
    absent_days = 0
    from datetime import datetime as dt_parser
    for s in summaries:
        if not s.get("is_present", False):
            try:
                d_obj = dt_parser.strptime(s["date"], "%Y-%m-%d")
                d_date = d_obj.date()
                if d_date in daily_schedules_map:
                    is_workday = not daily_schedules_map[d_date].is_off
                else:
                    is_workday = (d_obj.weekday() + 1) in work_days
            except Exception:
                is_workday = True
            
            if is_workday:
                absent_days += 1

    late_days = sum(1 for s in summaries if s.get("is_present", False) and s.get("is_late", False))
    incomplete_days = sum(1 for s in summaries if s.get("is_present", False) and s.get("missing_punches", False))
    
    punctuality_rate = "100%"
    if present_days > 0:
        ontime_days = present_days - late_days
        punctuality_rate = f"{round((ontime_days / present_days) * 100, 1)}%"

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        "BannerTitle_" + str(employee_code),
        parent=styles["Title"],
        fontSize=14,
        leading=16,
        textColor=colors.white,
        alignment=0,
        fontName="Helvetica-Bold"
    )
    
    banner_meta_style = ParagraphStyle(
        "BannerMeta_" + str(employee_code),
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#cbd5e1"),
        alignment=0
    )
    
    section_title = ParagraphStyle(
        "SecTitle_" + str(employee_code),
        parent=styles["Heading2"],
        fontSize=10,
        leading=13,
        textColor=primary_color,
        fontName="Helvetica-Bold",
        spaceBefore=6,
        spaceAfter=3
    )

    kpi_num_style = ParagraphStyle(
        "KpiNum_" + str(employee_code),
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=15,
        textColor=primary_color,
        alignment=1
    )

    kpi_lbl_style = ParagraphStyle(
        "KpiLbl_" + str(employee_code),
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#475569"),
        alignment=1
    )

    cell_style = ParagraphStyle(
        "CellNormal_" + str(employee_code),
        fontName="Helvetica",
        fontSize=7.0,
        leading=8.5,
        textColor=colors.HexColor("#334155"),
        alignment=1
    )

    cell_bold_style = ParagraphStyle(
        "CellBold_" + str(employee_code),
        parent=cell_style,
        fontName="Helvetica-Bold"
    )

    elements = []

    # 1. HEADER BANNER
    left_cell_elements = []
    has_logo = False
    if cfg.get("logo_path"):
        try:
            clean_path = cfg["logo_path"].split('?')[0]
            if clean_path.startswith('/'):
                clean_path = clean_path[1:]
            local_path = os.path.abspath(clean_path)
            if not os.path.exists(local_path) and not clean_path.startswith("uploads/"):
                local_path = os.path.abspath(os.path.join("uploads", clean_path))
                
            if os.path.exists(local_path):
                logo_img = CircularAvatar(local_path, 1.3*cm)
                logo_img.hAlign = 'LEFT'
                left_cell_elements.append(logo_img)
                left_cell_elements.append(Spacer(1, 4))
                has_logo = True
        except Exception:
            pass

    sys_title_style = ParagraphStyle(
        "SysTitle_" + str(employee_code),
        parent=title_style,
        fontSize=12 if has_logo else 15,
        leading=14 if has_logo else 17
    )
    left_cell_elements.append(Paragraph(cfg.get("system_name", "Control de Asistencia").upper(), sys_title_style))

    photo_img = None
    if photo_path:
        try:
            clean_photo_path = photo_path.split('?')[0]
            if clean_photo_path.startswith('/'):
                clean_photo_path = clean_photo_path[1:]
            local_photo_path = os.path.abspath(clean_photo_path)
            if not os.path.exists(local_photo_path) and not clean_photo_path.startswith("uploads/"):
                local_photo_path = os.path.abspath(os.path.join("uploads", clean_photo_path))
                
            if os.path.exists(local_photo_path):
                photo_img = CircularAvatar(local_photo_path, 2.2*cm)
        except Exception as pe:
            print(f"Error loading photo for PDF: {pe}")
            
    if not photo_img:
        photo_img = CircularAvatar(None, 2.2*cm)
        
    photo_img.hAlign = 'RIGHT'

    banner_data = [
        [
            left_cell_elements,
            Paragraph(f"<b>Empresa:</b> {cfg.get('company_name')}<br/><b>Empleado:</b> {employee_name}<br/><b>Código:</b> {employee_code}<br/><b>Área:</b> {department}", banner_meta_style),
            photo_img
        ]
    ]
    banner_table = Table(banner_data, colWidths=[7.8*cm, 8.5*cm, 2.5*cm])
    banner_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), primary_color),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (2, 0), (2, 0), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 16),
        ("RIGHTPADDING", (0, 0), (-1, -1), 16),
        ("BOTTOMBORDER", (0, 0), (-1, -1), 3, accent_color),
    ]))
    elements.append(banner_table)
    elements.append(Spacer(1, 0.15*cm))

    # Info de generación
    gen_time_str = get_local_now().strftime('%d/%m/%Y %H:%M')
    elements.append(Paragraph(f"<font color='#64748b'>Reporte generado en: {gen_time_str} — Periodo de análisis: {total_days} días analizados (Ausencias calculadas sobre días laborales laborados)</font>", banner_meta_style))
    elements.append(Spacer(1, 0.2*cm))

    # 2. KPIs METRICS CARDS
    kpis_data = [
        [
            Paragraph(punctuality_rate, kpi_num_style),
            Paragraph(str(present_days), kpi_num_style),
            Paragraph(str(absent_days), kpi_num_style),
            Paragraph(str(incomplete_days), kpi_num_style)
        ],
        [
            Paragraph("Tasa Puntualidad", kpi_lbl_style),
            Paragraph("Días Asistidos", kpi_lbl_style),
            Paragraph("Ausencias (Laborales)", kpi_lbl_style),
            Paragraph("Punches Incompletos", kpi_lbl_style)
        ]
    ]
    kpi_table = Table(kpis_data, colWidths=[4.7*cm, 4.7*cm, 4.7*cm, 4.7*cm])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fafafa")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#e2e8f0")),
        ("LINEABOVE", (0, 0), (-1, 0), 2.5, primary_color),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
        ("TOPPADDING", (0, 1), (-1, 1), 2),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
    ]))
    elements.append(kpi_table)

    # 3. CHARTS CONTAINER
    try:
        chart_doughnut_buf, chart_trend_buf = _generate_matplotlib_charts(summaries, granularity)
        img_doughnut = Image(chart_doughnut_buf, width=3.8*cm, height=3.8*cm)
        img_trend = Image(chart_trend_buf, width=7.6*cm, height=3.8*cm)
        
        charts_data = [[img_doughnut, img_trend]]
        charts_table = Table(charts_data, colWidths=[4.3*cm, 14.5*cm])
        charts_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        
        elements.append(Paragraph("ANÁLISIS GRÁFICO DE RENDIMIENTO", section_title))
        elements.append(charts_table)
    except Exception as e:
        elements.append(Paragraph(f"<font color='red'>No se pudieron cargar los gráficos: {str(e)}</font>", banner_meta_style))

    # 4. DETAILED DAILY ATTENDANCE TABLE
    elements.append(Paragraph("HISTORIAL DETALLADO DE ASISTENCIA", section_title))
    
    if granularity == "daily":
        all_cols = [
            ("period", "Fecha", 2.6*cm),
            ("schedule", "Horario", 2.5*cm),
            ("punches", "Entrada", 2.5*cm),
            ("punches", "Sal. Alm.", 2.5*cm),
            ("punches", "Ret. Alm.", 2.5*cm),
            ("punches", "Salida", 2.5*cm),
            ("status", "Estado", 3.5*cm)
        ]
    else:
        all_cols = [
            ("period", "Período", 3.6*cm),
            ("schedule", "Horario", 3.0*cm),
            ("status", "Asistió", 3.0*cm),
            ("status", "Tardanza", 3.0*cm),
            ("status", "Incompleto", 3.0*cm),
            ("punches", "Eventos", 3.0*cm)
        ]

    active_cols = []
    for col_key, col_header, col_w in all_cols:
        if columns is None or col_key in columns:
            active_cols.append((col_key, col_header, col_w))

    headers = [Paragraph(f"<b>{h}</b>", ParagraphStyle("H_" + h.replace(".", "").replace(" ", "_") + "_" + str(employee_code), parent=cell_bold_style, textColor=colors.white)) for _, h, _ in active_cols]
    data = [headers]
    
    total_w = sum(w for _, _, w in active_cols)
    if total_w > 0:
        scale = (18.8 * cm) / total_w
        col_widths = [w * scale for _, _, w in active_cols]
    else:
        col_widths = []

    if granularity == "daily":
        for s in summaries:
            def format_time(iso_str):
                if not iso_str:
                    return "-"
                try:
                    return datetime.fromisoformat(iso_str).strftime("%H:%M")
                except:
                    return "-"

            is_split = s.get("schedule_type") == "split"
            entry_1 = format_time(s["punches"].get("entry_1"))
            exit_1 = format_time(s["punches"].get("exit_1")) if is_split else "-"
            entry_2 = format_time(s["punches"].get("entry_2")) if is_split else "-"
            exit_2 = format_time(s["punches"].get("exit_2")) if is_split else format_time(s["punches"].get("exit_1"))

            leave_map = {
                "vacation": ("Vacaciones", "#0f766e"),
                "medical": ("Incapacidad", "#7c3aed"),
                "paid_leave": ("Licencia Rem.", "#2563eb"),
                "unpaid_leave": ("Licencia No Rem.", "#4f46e5"),
                "suspension": ("Suspensión", "#b91c1c")
            }
            status_text = "OK"
            status_color = "#166534"
            if s.get("justification"):
                status_color = "#7c3aed"  # Accent/Purple color for justifications
                if s["justification"]["override_status"] == "present":
                    status_text = "Justificado"
                else:
                    status_text = "A Tiempo (Just.)"
            elif s.get("leave_type") and s["leave_type"] in leave_map:
                status_text, status_color = leave_map[s["leave_type"]]
            elif not s.get("is_present", False):
                is_workday = True
                try:
                    d_obj = datetime.strptime(s["date"], "%Y-%m-%d")
                    d_date = d_obj.date()
                    if d_date in daily_schedules_map:
                        is_workday = not daily_schedules_map[d_date].is_off
                    else:
                        is_workday = (d_obj.weekday() + 1) in work_days
                except Exception:
                    pass
                
                if is_workday:
                    status_text = "Ausente"
                    status_color = "#991b1b"
                else:
                    status_text = "Descanso"
                    status_color = "#475569"
            elif s.get("missing_punches", False):
                status_text = "Incompleto"
                status_color = "#854d0e"
            elif s.get("is_late", False):
                status_text = "Tardanza"
                status_color = "#9a3412"

            status_paragraph = Paragraph(f"<b><font color='{status_color}'>{status_text}</font></b>", cell_bold_style)

            try:
                formatted_date = datetime.strptime(s["date"], "%Y-%m-%d").strftime("%d/%m/%Y")
            except:
                formatted_date = s.get("date", "")

            full_row_data = {
                "period": Paragraph(formatted_date, cell_style),
                "schedule": Paragraph("Partido" if is_split else ("Continuo" if s.get("schedule_type") == "continuous" else "Sin Horario"), cell_style),
                "punches_e1": Paragraph(entry_1, cell_style),
                "punches_x1": Paragraph(exit_1, cell_style),
                "punches_e2": Paragraph(entry_2, cell_style),
                "punches_x2": Paragraph(exit_2, cell_style),
                "status": status_paragraph
            }
            row = []
            for col_key, col_header, _ in active_cols:
                if col_header == "Entrada":
                    row.append(full_row_data["punches_e1"])
                elif col_header == "Sal. Alm.":
                    row.append(full_row_data["punches_x1"])
                elif col_header == "Ret. Alm.":
                    row.append(full_row_data["punches_e2"])
                elif col_header == "Salida":
                    row.append(full_row_data["punches_x2"])
                else:
                    row.append(full_row_data[col_key])
            data.append(row)
    else:
        for s in summaries:
            try:
                p_start = datetime.strptime(s["period_start"], "%Y-%m-%d").strftime("%d/%m/%Y")
                p_end = datetime.strptime(s["period_end"], "%Y-%m-%d").strftime("%d/%m/%Y")
                period_str = f"{p_start} - {p_end}"
            except:
                period_str = f"{s.get('period_start', '')} - {s.get('period_end', '')}"

            full_row_data = {
                "period": Paragraph(period_str, cell_style),
                "schedule": Paragraph("Partido" if s.get("schedule_type") == "split" else ("Continuo" if s.get("schedule_type") == "continuous" else "Sin Horario"), cell_style),
                "status_p": Paragraph("Sí" if s.get("is_present") else "No", ParagraphStyle("B_P_" + str(employee_code), parent=cell_bold_style, textColor=colors.HexColor("#00E676") if s.get("is_present") else colors.HexColor("#FF3D00"))),
                "status_l": Paragraph("Sí" if s.get("is_late") else "No", ParagraphStyle("B_T_" + str(employee_code), parent=cell_bold_style, textColor=colors.HexColor("#FFB300") if s.get("is_late") else colors.HexColor("#00E676"))),
                "status_m": Paragraph("Sí" if s.get("missing_punches") else "No", ParagraphStyle("B_I_" + str(employee_code), parent=cell_bold_style, textColor=colors.HexColor("#FFA000") if s.get("missing_punches") else colors.HexColor("#00E676"))),
                "punches": Paragraph(str(s.get("total_raw_events", 0)), cell_style)
            }
            row = []
            for col_key, col_header, _ in active_cols:
                if col_header == "Asistió":
                    row.append(full_row_data["status_p"])
                elif col_header == "Tardanza":
                    row.append(full_row_data["status_l"])
                elif col_header == "Incompleto":
                    row.append(full_row_data["status_m"])
                else:
                    row.append(full_row_data[col_key])
            data.append(row)

    if col_widths:
        table = Table(data, colWidths=col_widths, repeatRows=1)
        t_style = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), primary_color),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ])

        if granularity == "daily":
            for r in range(1, len(data)):
                s = summaries[r - 1]
                is_present = s.get("is_present", False)
                missing = s.get("missing_punches", False)
                is_late = s.get("is_late", False)
                
                if not is_present:
                    is_workday = True
                    try:
                        d_obj = datetime.strptime(s["date"], "%Y-%m-%d")
                        d_date = d_obj.date()
                        if d_date in daily_schedules_map:
                            is_workday = not daily_schedules_map[d_date].is_off
                        else:
                            is_workday = (d_obj.weekday() + 1) in work_days
                    except Exception:
                        pass
                    
                    if is_workday:
                        bg_color = colors.HexColor("#fee2e2") # soft red
                    else:
                        bg_color = colors.HexColor("#f1f5f9") # soft gray/blue
                elif missing:
                    bg_color = colors.HexColor("#fef9c3") # soft yellow
                elif is_late:
                    bg_color = colors.HexColor("#fef3c7") # soft orange
                else:
                    bg_color = colors.HexColor("#dcfce7") # soft green
                    
                t_style.add("BACKGROUND", (0, r), (-1, r), bg_color)
        else:
            for r in range(1, len(data)):
                bg = colors.white if r % 2 == 1 else colors.HexColor("#f8fafc")
                t_style.add("BACKGROUND", (0, r), (-1, r), bg)
                
                s = summaries[r - 1]
                for col_idx, (col_key, col_header, _) in enumerate(active_cols):
                    if col_header == "Asistió" and not s.get("is_present"):
                        t_style.add("BACKGROUND", (col_idx, r), (col_idx, r), colors.HexColor("#fee2e2"))
                    elif col_header == "Tardanza" and s.get("is_late"):
                        t_style.add("BACKGROUND", (col_idx, r), (col_idx, r), colors.HexColor("#fef3c7"))
                    elif col_header == "Incompleto" and s.get("missing_punches"):
                        t_style.add("BACKGROUND", (col_idx, r), (col_idx, r), colors.HexColor("#fef9c3"))

        table.setStyle(t_style)
        elements.append(table)

    return elements


def _generate_individual_pdf(summaries: list, granularity: str, columns: Optional[list] = None) -> bytes:
    """Genera un reporte PDF con diseño ejecutivo vertical para un empleado individual."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.lib.units import cm

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=0.8*cm, bottomMargin=1.6*cm, leftMargin=0.8*cm, rightMargin=0.8*cm
    )
    elements = _get_employee_flowables(summaries, granularity, columns=columns)
    doc.build(elements, canvasmaker=NumberedCanvas)
    return doc.filename if hasattr(doc, 'filename') else buf.getvalue()


def _generate_grouped_pdf(summaries: list, granularity: str, progress_callback = None, columns: Optional[list] = None) -> bytes:
    """Genera un reporte PDF unificado con diseño ejecutivo vertical por persona separado por PageBreaks."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, PageBreak, Paragraph, Spacer
    from reportlab.lib.units import cm
    from collections import defaultdict

    if not summaries:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            topMargin=2*cm, bottomMargin=2*cm, leftMargin=2*cm, rightMargin=2*cm
        )
        styles = getSampleStyleSheet()
        elements = [
            Paragraph("Reporte de Asistencia Vacío", ParagraphStyle("TitleEmpty", parent=styles["Title"], textColor=colors.HexColor("#1e3a5f"))),
            Spacer(1, 1*cm),
            Paragraph("No se encontraron registros de asistencia para los filtros y fechas seleccionadas.", ParagraphStyle("BodyEmpty", parent=styles["Normal"], textColor=colors.HexColor("#64748b")))
        ]
        doc.build(elements)
        if progress_callback:
            try:
                progress_callback(100)
            except:
                pass
        return buf.getvalue()

    # Precargar horarios de empleados de la BD en una sola consulta para evitar N+1 queries
    from backend.database import SessionLocal
    from backend.models import Employee
    from sqlalchemy.orm import joinedload
    
    schedules_map = {}
    db = SessionLocal()
    try:
        employees = db.query(Employee).options(joinedload(Employee.schedule)).all()
        for emp in employees:
            if emp.schedule and emp.schedule.work_days:
                schedules_map[emp.id] = [int(x) for x in emp.schedule.work_days.split(",")]
    except Exception as e:
        print(f"Error precargando horarios de empleados en reporte PDF: {e}")
    finally:
        db.close()

    # Agrupar por ID de empleado
    grouped = defaultdict(list)
    for s in summaries:
        emp_id = s.get("employee_id")
        if emp_id is not None:
            grouped[emp_id].append(s)

    # Ordenar empleados por nombre alfabéticamente
    sorted_emp_ids = sorted(
        grouped.keys(),
        key=lambda eid: grouped[eid][0].get("employee_name", "").lower()
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=0.8*cm, bottomMargin=1.6*cm, leftMargin=0.8*cm, rightMargin=0.8*cm
    )
    
    elements = []
    total_emps = len(sorted_emp_ids)
    for idx, emp_id in enumerate(sorted_emp_ids):
        emp_summaries = grouped[emp_id]
        emp_summaries = sorted(emp_summaries, key=lambda x: x.get("date", ""))
        
        emp_flowables = _get_employee_flowables(emp_summaries, granularity, schedules_map=schedules_map, columns=columns)
        elements.extend(emp_flowables)
        
        # Añadir salto de página si no es el último empleado
        if idx < total_emps - 1:
            elements.append(PageBreak())
            
        if progress_callback:
            try:
                progress_callback(int((idx + 1) / total_emps * 100))
            except Exception as pe:
                print(f"Error en progress_callback: {pe}")
            
    doc.build(elements, canvasmaker=NumberedCanvas)
    return buf.getvalue()


def generate_attendance_pdf(summaries: list, granularity: str, progress_callback = None, columns: Optional[list] = None) -> bytes:
    """Genera un reporte de asistencia unificado en PDF (agrupado por persona si hay varios)."""
    return _generate_grouped_pdf(summaries, granularity, progress_callback=progress_callback, columns=columns)

