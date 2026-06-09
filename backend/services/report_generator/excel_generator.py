"""
Generador de reportes Excel.
"""
import io
from datetime import datetime
from typing import Optional
from sqlalchemy.orm import Session
from backend.models import AttendanceRecord, Employee

def generate_excel_report(
    db: Session,
    employee_id: Optional[int] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    department_id: Optional[int] = None,
    position_id: Optional[int] = None,
    schedule_id: Optional[int] = None,
) -> bytes:
    """Genera un reporte Excel con los registros de asistencia filtrados."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    query = db.query(AttendanceRecord).join(
        Employee, AttendanceRecord.employee_id == Employee.id, isouter=True
    )
    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if position_id:
        query = query.filter(Employee.position_id == position_id)
    if schedule_id:
        query = query.filter(Employee.schedule_id == schedule_id)
    if date_from:
        query = query.filter(AttendanceRecord.event_time >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.event_time <= date_to)

    records = query.order_by(AttendanceRecord.event_time.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "Empleado", "Código", "Departamento", "Fecha", "Hora", "Tipo", "Método Auth", "Tardanza"]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Datos
    for i, r in enumerate(records, 1):
        emp = r.employee
        row = [
            i,
            emp.full_name if emp else "Desconocido",
            emp.employee_code if emp else "-",
            emp.department.name if emp and emp.department else "-",
            r.event_time.strftime("%d/%m/%Y"),
            r.event_time.strftime("%H:%M:%S"),
            "Entrada" if r.event_type == "entry" else "Salida",
            r.auth_method or "-",
            "Sí" if r.is_late else "No",
        ]
        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.alignment = center
            cell.border = border
            if r.is_late and col_idx == 9:
                cell.font = Font(color="CC0000", bold=True)

    # Anchos de columna
    col_widths = [5, 28, 12, 20, 14, 12, 10, 16, 10]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generate_consolidated_excel(
    db: Session,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    department_id: Optional[int] = None,
    position_id: Optional[int] = None,
    schedule_id: Optional[int] = None,
    columns: Optional[list] = None,
) -> bytes:
    """Genera un reporte consolidado en Excel con entrada y salida en la misma fila por día."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    query = db.query(AttendanceRecord).join(
        Employee, AttendanceRecord.employee_id == Employee.id, isouter=True
    )
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if position_id:
        query = query.filter(Employee.position_id == position_id)
    if schedule_id:
        query = query.filter(Employee.schedule_id == schedule_id)
    if date_from:
        query = query.filter(AttendanceRecord.event_time >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.event_time <= date_to)

    records = query.order_by(AttendanceRecord.event_time.asc()).all()

    # Agrupar por (empleado, fecha)
    data_map = defaultdict(lambda: {"entries": [], "exits": []})
    employee_cache = {}

    for r in records:
        if not r.employee_id:
            continue
        emp_id = r.employee_id
        if emp_id not in employee_cache:
            employee_cache[emp_id] = r.employee
            
        day = r.event_time.date()
        if r.event_type == "entry":
            data_map[(emp_id, day)]["entries"].append(r)
        else:
            data_map[(emp_id, day)]["exits"].append(r)

    # Ordenar por fecha desc, luego nombre empleado
    sorted_keys = sorted(
        data_map.keys(),
        key=lambda k: (k[1], employee_cache.get(k[0]).last_name if employee_cache.get(k[0]) else ""),
        reverse=True
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado Diario"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    all_cols = [
        ("period", "Fecha"),
        ("employee_code", "Código"),
        ("employee_name", "Empleado"),
        ("department", "Departamento"),
        ("punches", "Hora Entrada"),
        ("punches", "Hora Salida"),
        ("schedule", "Horario"),
        ("status", "Tardanza")
    ]

    active_cols = []
    for col_key, col_header in all_cols:
        if columns is None or col_key in columns:
            active_cols.append((col_key, col_header))

    headers = [h for _, h in active_cols]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    row_num = 2
    for key in sorted_keys:
        emp_id, day = key
        emp = employee_cache.get(emp_id)
        if not emp:
            continue
            
        group = data_map[key]
        
        # Entrada: primer registro de entrada del día
        entry_time_str = "-"
        is_late_str = "No"
        if group["entries"]:
            first_entry = min(group["entries"], key=lambda r: r.event_time)
            entry_time_str = first_entry.event_time.strftime("%H:%M:%S")
            is_late_str = "Sí" if first_entry.is_late else "No"
            
        # Salida: último registro de salida del día
        exit_time_str = "-"
        if group["exits"]:
            last_exit = max(group["exits"], key=lambda r: r.event_time)
            exit_time_str = last_exit.event_time.strftime("%H:%M:%S")

        full_row_data = {
            "period": day.strftime("%d/%m/%Y"),
            "employee_code": emp.employee_code,
            "employee_name": emp.full_name,
            "department": emp.department.name if emp.department else "-",
            "punches_in": entry_time_str,
            "punches_out": exit_time_str,
            "schedule": f"{emp.schedule.name} ({emp.schedule.work_start_time}-{emp.schedule.work_end_time})" if emp.schedule_id and emp.schedule else f"{emp.work_start_time} - {emp.work_end_time}",
            "status": is_late_str
        }

        row = []
        for col_key, col_header in active_cols:
            if col_header == "Hora Entrada":
                row.append(full_row_data["punches_in"])
            elif col_header == "Hora Salida":
                row.append(full_row_data["punches_out"])
            else:
                row.append(full_row_data[col_key])

        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = center
            cell.border = border
            header_name = headers[col_idx - 1]
            if header_name == "Tardanza" and is_late_str == "Sí":
                cell.font = Font(color="CC0000", bold=True)
        
        row_num += 1

    # Anchos de columna
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generate_attendance_excel(summaries: list, granularity: str, columns: Optional[list] = None) -> bytes:
    """Genera un reporte consolidado con horarios y granularidad en Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {granularity.capitalize()}"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if granularity == "daily":
        all_cols = [
            ("employee_name", "Empleado"),
            ("employee_code", "Código"),
            ("department", "Departamento"),
            ("period", "Fecha"),
            ("schedule", "Horario"),
            ("punches", "Entrada"),
            ("punches", "Salida Almuerzo"),
            ("punches", "Retorno Almuerzo"),
            ("punches", "Salida"),
            ("status", "Estado")
        ]
    else:
        all_cols = [
            ("employee_name", "Empleado"),
            ("employee_code", "Código"),
            ("department", "Departamento"),
            ("period", "Período"),
            ("schedule", "Horario"),
            ("status", "Asistente"),
            ("status", "Incidencias Tardanza"),
            ("status", "Punches Incompletos"),
            ("punches", "Eventos Totales")
        ]

    active_cols = []
    for col_key, col_header in all_cols:
        if columns is None or col_key in columns:
            active_cols.append((col_key, col_header))

    headers = [h for _, h in active_cols]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    row_num = 2
    for s in summaries:
        if granularity == "daily":
            def format_time(iso_str):
                if not iso_str:
                    return "-"
                try:
                    return datetime.fromisoformat(iso_str).strftime("%H:%M")
                except:
                    return "-"

            is_split = s["schedule_type"] == "split"
            entry_1 = format_time(s["punches"]["entry_1"])
            exit_1 = format_time(s["punches"]["exit_1"]) if is_split else "-"
            entry_2 = format_time(s["punches"]["entry_2"]) if is_split else "-"
            exit_2 = format_time(s["punches"]["exit_2"]) if is_split else format_time(s["punches"]["exit_1"])

            leave_map = {
                "vacation": "Vacaciones",
                "medical": "Incapacidad",
                "paid_leave": "Licencia Remunerada",
                "unpaid_leave": "Licencia No Remunerada",
                "suspension": "Suspensión"
            }
            status = "OK"
            if s.get("justification"):
                if s["justification"]["override_status"] == "present":
                    status = "Justificado"
                else:
                    status = "A Tiempo (Just.)"
            elif s.get("leave_type") and s["leave_type"] in leave_map:
                status = leave_map[s["leave_type"]]
            elif not s["is_present"]:
                if s.get("is_holiday"):
                    status = "Festivo"
                else:
                    status = "Ausente"
            elif s["missing_punches"]:
                status = "Incompleto"
            elif s["is_late"]:
                status = "Tardanza"

            try:
                formatted_date = datetime.strptime(s["date"], "%Y-%m-%d").strftime("%d/%m/%Y")
            except:
                formatted_date = s["date"]

            full_row_data = {
                "employee_name": s["employee_name"],
                "employee_code": s["employee_code"],
                "department": s["department"],
                "period": formatted_date,
                "schedule": "Partido" if is_split else ("Continuo" if s["schedule_type"] == "continuous" else "Sin Horario"),
                "punches_e1": entry_1,
                "punches_x1": exit_1,
                "punches_e2": entry_2,
                "punches_x2": exit_2,
                "status": status
            }

            row = []
            for col_key, col_header in active_cols:
                if col_header == "Entrada":
                    row.append(full_row_data["punches_e1"])
                elif col_header == "Salida Almuerzo":
                    row.append(full_row_data["punches_x1"])
                elif col_header == "Retorno Almuerzo":
                    row.append(full_row_data["punches_e2"])
                elif col_header == "Salida":
                    row.append(full_row_data["punches_x2"])
                else:
                    row.append(full_row_data[col_key])
        else:
            try:
                p_start = datetime.strptime(s["period_start"], "%Y-%m-%d").strftime("%d/%m/%Y")
                p_end = datetime.strptime(s["period_end"], "%Y-%m-%d").strftime("%d/%m/%Y")
                period_str = f"{p_start} - {p_end}"
            except:
                period_str = f"{s['period_start']} - {s['period_end']}"

            full_row_data = {
                "employee_name": s["employee_name"],
                "employee_code": s["employee_code"],
                "department": s["department"],
                "period": period_str,
                "schedule": "Partido" if s["schedule_type"] == "split" else ("Continuo" if s["schedule_type"] == "continuous" else "Sin Horario"),
                "status_p": "Sí" if s["is_present"] else "No",
                "status_l": "Sí" if s["is_late"] else "No",
                "status_m": "Sí" if s["missing_punches"] else "No",
                "punches": s["total_raw_events"]
            }

            row = []
            for col_key, col_header in active_cols:
                if col_header == "Asistente":
                    row.append(full_row_data["status_p"])
                elif col_header == "Incidencias Tardanza":
                    row.append(full_row_data["status_l"])
                elif col_header == "Punches Incompletos":
                    row.append(full_row_data["status_m"])
                else:
                    row.append(full_row_data[col_key])

        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = center
            cell.border = border
            header_name = headers[col_idx - 1]
            if granularity == "daily" and header_name == "Estado":
                if status == "Ausente":
                    cell.font = Font(color="FF3D00", bold=True)
                elif status == "Tardanza":
                    cell.font = Font(color="FFB300", bold=True)
                elif status == "Incompleto":
                    cell.font = Font(color="FFA000", bold=True)
                else:
                    cell.font = Font(color="00E676", bold=True)
            elif granularity != "daily":
                if header_name == "Incidencias Tardanza" and s["is_late"]:
                    cell.font = Font(color="FFB300", bold=True)
                elif header_name == "Punches Incompletos" and s["missing_punches"]:
                    cell.font = Font(color="FFA000", bold=True)

        row_num += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
